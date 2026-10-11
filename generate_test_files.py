import os
from pathlib import Path
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
import pandas as pd
from docx import Document

def ensure_dir(p):
    os.makedirs(p, exist_ok=True)

base = Path('test_files')
ensure_dir(base)
# txt
(base / 'sample.txt').write_text('Hello world! This is a test txt file.')
# csv
pd.DataFrame({'col1':[1,2,3], 'col2':['a','b','c']}).to_csv(base/'sample.csv', index=False)
# xlsx
pd.DataFrame({'num':[10,20,30], 'val':['x','y','z']}).to_excel(base/'sample.xlsx', index=False)
# docx
doc = Document()
doc.add_heading('Sample Docx', level=1)
doc.add_paragraph('This is a paragraph in the docx file.')
doc.save(base/'sample.docx')
# pdf
c = canvas.Canvas(str(base/'sample.pdf'), pagesize=letter)
c.drawString(100, 750, 'Hello PDF! This is a test PDF file.')
c.save()
# contradictory files
(base / 'conflict1.txt').write_text('The critical value is 10.')
(base / 'conflict2.txt').write_text('The critical value is 20.')
print('Test files generated with contradictions.')
